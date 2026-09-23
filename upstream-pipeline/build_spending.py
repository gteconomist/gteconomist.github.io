import openpyxl, collections
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L

SRC = 'GT_AY2025_Enrollment_Geographic_Summary.xlsx'
OUT = 'GT_AY2025_Student_Spending_Estimate.xlsx'

# ---------- enrollment cross-tab from the Data tab ----------
wb0 = openpyxl.load_workbook(SRC, data_only=True)
rows = list(wb0['Data'].iter_rows(values_only=True))[1:]
xw = collections.Counter()
for ay, lvl, dl, hs, src, gc, reg, cty, n in rows:
    if gc == 'Georgia':
        geo = 'GA - Atlanta region (ARC)' if reg == 'Atlanta Regional Commission' else 'GA - other regions'
    elif gc in ('Other U.S. State', 'U.S. Territory', 'Military / APO-FPO'):
        geo = 'Other U.S. (incl. territories/military)'
    else:
        geo = 'Not reported (in-person = international)'
    xw[(lvl, dl, hs, geo)] += int(n)
LEVELS = ['Undergraduate', 'Graduate']; DELIV = ['In-Person', 'Online']; HOUS = ['Yes', 'No']
GEOS = ['GA - Atlanta region (ARC)', 'GA - other regions', 'Other U.S. (incl. territories/military)', 'Not reported (in-person = international)']
RATEGEO = {GEOS[0]: 'Georgia resident', GEOS[1]: 'Georgia resident', GEOS[2]: 'Out of state', GEOS[3]: 'International'}

# ---------- styles ----------
F = 'Arial'
def font(bold=False, color='000000', size=10, italic=False): return Font(name=F, bold=bold, color=color, size=size, italic=italic)
BLUE = '0000FF'; GREEN = '008000'
YELLOW = PatternFill('solid', fgColor='FFFF00'); HEAD = PatternFill('solid', fgColor='003057'); SUB = PatternFill('solid', fgColor='DCE6F1'); TOT = PatternFill('solid', fgColor='EEEEEE')
thin = Side(style='thin', color='BBBBBB'); BOX = Border(top=thin, bottom=thin, left=thin, right=thin)
USD = '$#,##0;($#,##0);-'; NUM = '#,##0;(#,##0);-'; PCT = '0.0%'

wb = openpyxl.Workbook()

def header(ws, row, labels, col=1):
    for i, l in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=l); c.font = Font(name=F, bold=True, color='FFFFFF', size=10); c.fill = HEAD
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True); c.border = BOX
    ws.row_dimensions[row].height = 30

def title(ws, text, sub=None):
    ws['A1'] = text; ws['A1'].font = font(True, '003057', 14)
    if sub: ws['A2'] = sub; ws['A2'].font = font(italic=True, color='555555', size=9)

# ================= Assumptions =================
wa = wb.active; wa.title = 'Assumptions'
title(wa, 'Student Spending Estimate — Assumptions', 'Yellow cells are inputs (blue text). Everything else in the workbook recalculates from these.')
A = [
 ('Freshman headcount (no housing allowance; required to live on campus)', 3838, NUM, 'GT fall 2024 first-time full-time class. Source: gatech.edu/news 2024-11-12 "Georgia Tech Reaches New Milestones…". Allocated across geographies in proportion to in-person undergrads in on-campus housing.'),
 ('Share of ARC-region undergrads NOT in campus housing who live at home', 0.5, PCT, 'Assumption (Alfie, 2026-09-21). Enrollment file cannot distinguish off-campus from living-at-home. Non-ARC Georgians and out-of-state students are treated as living off campus.'),
 ('Share of ARC-region graduate students NOT in campus housing who live at home', 0.0, PCT, 'Assumption (Alfie, 2026-09-21). If >0, the UG "current/transfer living with parent" budget is applied (no graduate at-home budget is published).'),
 ('Include on-campus housing allowance for NON-freshmen in campus housing? (1 = yes, 0 = no)', 1, '0', 'Alfie 2026-09-21: include. Set to 0 to treat it as Institute revenue already counted elsewhere (same logic as the freshman exclusion).'),
 ('Spending factor for online / distance students (0 = no spending)', 0.0, PCT, 'Alfie 2026-09-21: assume no spending for now. If >0, the off-campus budget for their level/geography is applied times this factor.'),
 ('Treat in-person students with no reported county as international? (1 = yes, 0 = exclude)', 1, '0', 'Alfie 2026-09-21: yes. 5,634 in-person students have a blank county in the extract; the out-of-country budget is applied.'),
]
header(wa, 4, ['Assumption', 'Value', 'Notes / source'])
for i, (lab, val, fmt, note) in enumerate(A):
    r = 5 + i
    wa.cell(row=r, column=1, value=lab).font = font(); wa.cell(row=r, column=1).alignment = Alignment(wrap_text=True, vertical='top')
    c = wa.cell(row=r, column=2, value=val); c.font = font(color=BLUE); c.fill = YELLOW; c.number_format = fmt; c.border = BOX; c.alignment = Alignment(horizontal='center', vertical='top')
    wa.cell(row=r, column=3, value=note).font = font(size=9, color='555555'); wa.cell(row=r, column=3).alignment = Alignment(wrap_text=True, vertical='top')
    wa.row_dimensions[r].height = 42
wa.column_dimensions['A'].width = 62; wa.column_dimensions['B'].width = 12; wa.column_dimensions['C'].width = 90
FRESH, HOME_UG, HOME_GR, UPPER_HSG, ONLINE_F, NR_INTL = ['Assumptions!$B$%d' % (5 + i) for i in range(6)]
r = 12
wa.cell(row=r, column=1, value='Standing exclusions (per Alfie)').font = font(True)
for j, t in enumerate(['Tuition and mandatory student fees: excluded for every student — captured elsewhere in the Institute revenue figures.',
                       'Freshman housing allowance: excluded — freshmen must live on campus, and that revenue is captured elsewhere.',
                       'Average loan costs (graduate budgets, $140): excluded — a payment to lenders, not spending in the Georgia economy.',
                       'Online / distance students: no spending (factor above), pending a better assumption.']):
    wa.cell(row=r + 1 + j, column=1, value='• ' + t).font = font(size=9); wa.cell(row=r + 1 + j, column=1).alignment = Alignment(wrap_text=True)
    wa.row_dimensions[r + 1 + j].height = 26

# ================= Rates =================
wr = wb.create_sheet('Rates')
title(wr, 'Georgia Tech 2025-26 Cost of Attendance — published budgets and amounts applied',
      'Sources: finaid.gatech.edu/costs/undergraduate-costs and finaid.gatech.edu/costs/graduate-costs (accessed 2026-09-21). Per academic year (2 semesters). Applied = Published × Include.')
ITEMS = ['Tuition', 'Mandatory fees', 'Books & supplies', 'Housing', 'Food', 'Personal / misc.', 'Transportation', 'Loan costs']
SPEND = ['Books & supplies', 'Housing', 'Food', 'Personal / misc.', 'Transportation']
# published amounts by budget: dict geo -> list in ITEMS order
tu = {'Georgia resident': 10512, 'Out of state': 33596, 'International': 34572}
fe = {'Georgia resident': 1496, 'Out of state': 1496, 'International': 1696}
tr = {'Georgia resident': 550, 'Out of state': 950, 'International': 950}
def ug(geo, housing, food, transport=None):
    return [tu[geo], fe[geo], 800, housing, food, 2800, tr[geo] if transport is None else transport, 0]
def gr(geo, housing):
    return [0, 0, 800, housing, 6132, 2800, tr[geo], 140]
BUDGETS = [
 ('FY_ON', 'First-year, living on campus', lambda g: ug(g, 7864, 6132)),
 ('FY_OFF', 'First-year, living off campus', lambda g: ug(g, 11744, 6132)),
 ('FY_HOME', 'First-year, living with parent', lambda g: ug(g, 766, 3372, 3360)),
 ('CT_ON', 'Current/transfer UG, living on campus', lambda g: ug(g, 10528, 6132)),
 ('CT_OFF', 'Current/transfer UG, living off campus', lambda g: ug(g, 11744, 6132)),
 ('CT_HOME', 'Current/transfer UG, living with parent', lambda g: ug(g, 766, 3372, 3360)),
 ('GR_ON', 'Graduate, living on campus', lambda g: gr(g, 11532)),
 ('GR_OFF', 'Graduate, living off campus', lambda g: gr(g, 13974)),
]
# include flags per (budget, item)
def include(bkey, item):
    if item in ('Tuition', 'Mandatory fees', 'Loan costs'): return 0
    if item == 'Housing' and bkey == 'FY_ON': return 0
    if item == 'Housing' and bkey in ('CT_ON', 'GR_ON'): return '=' + UPPER_HSG
    return 1
header(wr, 4, ['Key', 'Budget', 'Geography'] + ['Published: ' + i for i in ITEMS] + ['Published total'] + ['Include: ' + i for i in ITEMS] + ['Applied: ' + i for i in ITEMS] + ['Applied total'])
GEOR = ['Georgia resident', 'Out of state', 'International']
r = 5; RATE_FIRST = r
for bkey, bname, fn in BUDGETS:
    for g in GEOR:
        vals = fn(g)
        wr.cell(row=r, column=1, value=f'{bkey}|{g}').font = font()
        wr.cell(row=r, column=2, value=bname).font = font(); wr.cell(row=r, column=3, value=g).font = font()
        for k, v in enumerate(vals):
            c = wr.cell(row=r, column=4 + k, value=v); c.font = font(color=BLUE); c.number_format = USD
        c = wr.cell(row=r, column=12, value=f'=SUM(D{r}:K{r})'); c.font = font(True); c.number_format = USD
        for k, it in enumerate(ITEMS):
            inc = include(bkey, it)
            c = wr.cell(row=r, column=13 + k, value=inc); c.number_format = '0'; c.alignment = Alignment(horizontal='center')
            c.font = font(color=GREEN) if isinstance(inc, str) else font(color=BLUE)
        for k in range(len(ITEMS)):
            c = wr.cell(row=r, column=21 + k, value=f'={L(4+k)}{r}*{L(13+k)}{r}'); c.number_format = USD; c.font = font()
        c = wr.cell(row=r, column=29, value=f'=SUM(U{r}:AB{r})'); c.font = font(True); c.number_format = USD
        r += 1
RATE_LAST = r - 1
r += 1
wr.cell(row=r, column=1, value='Notes').font = font(True)
for j, t in enumerate(['Published totals match the finaid.gatech.edu tables (e.g., first-year on campus GA resident $30,154; graduate off campus GA resident $24,396 excl. tuition/fees).',
                       'Graduate budgets on the finaid page exclude tuition and fees (listed at bursar.gatech.edu); they are excluded here anyway, so shown as 0.',
                       'Online/distance students are assigned the off-campus budget for their level and geography, multiplied by the Assumptions online factor (0 = no spending).',
                       'Include flags: 1 = counted as student spending, 0 = excluded. The non-freshman on-campus housing flag links to Assumptions.']):
    wr.cell(row=r + 1 + j, column=1, value='• ' + t).font = font(size=9)
wr.column_dimensions['A'].width = 26; wr.column_dimensions['B'].width = 38; wr.column_dimensions['C'].width = 16
for cidx in range(4, 30): wr.column_dimensions[L(cidx)].width = 12
wr.freeze_panes = 'D5'
RATE_KEY = f'Rates!$A${RATE_FIRST}:$A${RATE_LAST}'
def rate_col(item): return L(21 + ITEMS.index(item))
def RATE(item, keycell): return f'INDEX(Rates!${rate_col(item)}${RATE_FIRST}:${rate_col(item)}${RATE_LAST},MATCH({keycell},{RATE_KEY},0))'

# ================= Enrollment =================
we = wb.create_sheet('Enrollment')
title(we, 'AY2025 enrollment cross-tab — level × delivery × on-campus housing × geography',
      'Source: GT_AY2025_Enrollment_Geographic_Summary.xlsx (Data tab; 64,820 students). Territories and military addresses grouped with other U.S. states (out-of-state budget). Not-reported = blank county in the extract.')
header(we, 4, ['Level', 'Delivery', 'On-campus housing', 'Geography', 'Students'])
r = 5; EN_FIRST = r; ENR = {}
for lv in LEVELS:
    for dl in DELIV:
        for hs in HOUS:
            for g in GEOS:
                n = xw.get((lv, dl, hs, g), 0)
                for k, v in enumerate([lv, dl, hs, g]): we.cell(row=r, column=1 + k, value=v).font = font()
                c = we.cell(row=r, column=5, value=n); c.font = font(color=BLUE); c.number_format = NUM
                ENR[(lv, dl, hs, g)] = f'Enrollment!$E${r}'
                r += 1
EN_LAST = r - 1
c = we.cell(row=r, column=4, value='TOTAL'); c.font = font(True)
c = we.cell(row=r, column=5, value=f'=SUM(E{EN_FIRST}:E{EN_LAST})'); c.font = font(True); c.number_format = NUM; c.fill = TOT
for k, w in enumerate([16, 12, 18, 42, 12]): we.column_dimensions[L(1 + k)].width = w
we.freeze_panes = 'A5'

# ================= Allocation =================
wl = wb.create_sheet('Allocation')
title(wl, 'Allocation of students to cost-of-attendance budgets, and spending by category',
      'Student count formulas apply the Assumptions; spending = students × applied rate (Rates sheet). All dollars per academic year.')
header(wl, 4, ['Student category', 'Level', 'Geography (enrollment)', 'Rate geography', 'Budget key', 'Students'] + SPEND + ['Total spending', 'Per student'])
r = 5; AL_FIRST = r
def H(lv, g): return ENR[(lv, 'In-Person', 'Yes', g)]
def N(lv, g): return ENR[(lv, 'In-Person', 'No', g)]
def O(lv, g): return ENR[(lv, 'Online', 'Yes', g)] + '+' + ENR[(lv, 'Online', 'No', g)]
HSUM = '(' + '+'.join(H('Undergraduate', g) for g in GEOS) + ')'
def nr_gate(g): return f'*{NR_INTL}' if g == GEOS[3] else ''
CATS = []  # (category, level, geo, ratekey_budget, count_formula)
for g in GEOS:
    CATS.append(('First-year UG, on campus (no housing allowance)', 'Undergraduate', g, 'FY_ON', f'={FRESH}*{H("Undergraduate", g)}/{HSUM}{nr_gate(g)}'))
for g in GEOS:
    CATS.append(('Continuing UG, on campus', 'Undergraduate', g, 'CT_ON', f'=({H("Undergraduate", g)}-{FRESH}*{H("Undergraduate", g)}/{HSUM}){nr_gate(g)}'))
CATS.append(('Continuing UG, living at home', 'Undergraduate', GEOS[0], 'CT_HOME', f'={N("Undergraduate", GEOS[0])}*{HOME_UG}'))
for g in GEOS:
    home = f'*(1-{HOME_UG})' if g == GEOS[0] else ''
    CATS.append(('Continuing UG, off campus', 'Undergraduate', g, 'CT_OFF', f'={N("Undergraduate", g)}{home}{nr_gate(g)}'))
for g in GEOS:
    CATS.append(('Graduate, on campus', 'Graduate', g, 'GR_ON', f'={H("Graduate", g)}{nr_gate(g)}'))
CATS.append(('Graduate, living at home', 'Graduate', GEOS[0], 'CT_HOME', f'={N("Graduate", GEOS[0])}*{HOME_GR}'))
for g in GEOS:
    home = f'*(1-{HOME_GR})' if g == GEOS[0] else ''
    CATS.append(('Graduate, off campus', 'Graduate', g, 'GR_OFF', f'={N("Graduate", g)}{home}{nr_gate(g)}'))
for lv, bk in (('Undergraduate', 'CT_OFF'), ('Graduate', 'GR_OFF')):
    for g in GEOS:
        CATS.append((f'Online / distance ({lv.lower()}) × online factor', lv, g, bk, f'=({O(lv, g)})*{ONLINE_F}'))
for cat, lv, g, bk, cf in CATS:
    wl.cell(row=r, column=1, value=cat).font = font(); wl.cell(row=r, column=2, value=lv).font = font()
    wl.cell(row=r, column=3, value=g).font = font(); wl.cell(row=r, column=4, value=RATEGEO[g]).font = font()
    wl.cell(row=r, column=5, value=f'{bk}|{RATEGEO[g]}').font = font()
    c = wl.cell(row=r, column=6, value=cf); c.font = font(color=GREEN); c.number_format = NUM
    for k, it in enumerate(SPEND):
        c = wl.cell(row=r, column=7 + k, value=f'=$F{r}*{RATE(it, f"$E{r}")}'); c.font = font(); c.number_format = USD
    c = wl.cell(row=r, column=12, value=f'=SUM(G{r}:K{r})'); c.font = font(True); c.number_format = USD
    c = wl.cell(row=r, column=13, value=f'=IF(F{r}=0,0,L{r}/F{r})'); c.font = font(); c.number_format = USD
    r += 1
AL_LAST = r - 1
wl.cell(row=r, column=1, value='TOTAL').font = font(True)
for col in [6] + list(range(7, 13)):
    c = wl.cell(row=r, column=col, value=f'=SUM({L(col)}{AL_FIRST}:{L(col)}{AL_LAST})'); c.font = font(True); c.fill = TOT
    c.number_format = NUM if col == 6 else USD
c = wl.cell(row=r, column=13, value=f'=IF(F{r}=0,0,L{r}/F{r})'); c.font = font(True); c.number_format = USD; c.fill = TOT
AL_TOT = r
r += 1
c = wl.cell(row=r, column=1, value='Check: students allocated (excl. online rows, which are scaled by the online factor) vs. in-person enrollment'); c.font = font(size=9, italic=True)
c = wl.cell(row=r, column=6, value=f'=SUMIFS(F{AL_FIRST}:F{AL_LAST},B{AL_FIRST}:B{AL_LAST},"*",A{AL_FIRST}:A{AL_LAST},"<>Online*")-(SUMIFS(Enrollment!E{EN_FIRST}:E{EN_LAST},Enrollment!B{EN_FIRST}:B{EN_LAST},"In-Person")-{ENR[("Undergraduate","In-Person","Yes",GEOS[3])]}*(1-{NR_INTL})-{ENR[("Undergraduate","In-Person","No",GEOS[3])]}*(1-{NR_INTL})-{ENR[("Graduate","In-Person","Yes",GEOS[3])]}*(1-{NR_INTL})-{ENR[("Graduate","In-Person","No",GEOS[3])]}*(1-{NR_INTL}))')
c.number_format = NUM; c.font = font(size=9, italic=True)
wl.cell(row=r, column=7, value='← should be 0').font = font(size=9, italic=True)
for k, w in enumerate([44, 14, 40, 16, 26, 11, 14, 14, 14, 14, 14, 16, 12]): wl.column_dimensions[L(1 + k)].width = w
wl.freeze_panes = 'F5'

# ================= Summary =================
ws = wb.create_sheet('Summary', 0)
title(ws, 'Georgia Tech AY2025 Student Spending Estimate — for I-O model input',
      'Spending by category of student and category of spending, per academic year. Excludes tuition, mandatory fees, freshman housing, loan costs; online students at the Assumptions factor. Draft 2026-09-21.')
A_ = f'Allocation!$A${AL_FIRST}:$A${AL_LAST}'; B_ = f'Allocation!$B${AL_FIRST}:$B${AL_LAST}'; D_ = f'Allocation!$D${AL_FIRST}:$D${AL_LAST}'; C_ = f'Allocation!$C${AL_FIRST}:$C${AL_LAST}'
def col_(c): return f'Allocation!${c}${AL_FIRST}:${c}${AL_LAST}'
CATNAMES = []
for cat, *_ in CATS:
    if cat not in CATNAMES: CATNAMES.append(cat)

def block(r0, label, keys, crit_range, extra_crit=None):
    ws.cell(row=r0, column=1, value=label).font = font(True, '003057', 11)
    header(ws, r0 + 1, ['', 'Students'] + SPEND + ['Total spending', 'Share of total', 'Per student'])
    rr = r0 + 2
    for key in keys:
        ws.cell(row=rr, column=1, value=key).font = font()
        crit = f'{crit_range},"{key}"' + (f',{extra_crit}' if extra_crit else '')
        c = ws.cell(row=rr, column=2, value=f'=SUMIFS({col_("F")},{crit})'); c.number_format = NUM; c.font = font(color=GREEN)
        for k in range(len(SPEND)):
            c = ws.cell(row=rr, column=3 + k, value=f'=SUMIFS({col_(L(7+k))},{crit})'); c.number_format = USD; c.font = font(color=GREEN)
        c = ws.cell(row=rr, column=8, value=f'=SUM(C{rr}:G{rr})'); c.number_format = USD; c.font = font(True)
        rr += 1
    tr_ = rr
    ws.cell(row=tr_, column=1, value='Total').font = font(True)
    for col in range(2, 9):
        c = ws.cell(row=tr_, column=col, value=f'=SUM({L(col)}{r0+2}:{L(col)}{tr_-1})'); c.font = font(True); c.fill = TOT
        c.number_format = NUM if col == 2 else USD
    for x in range(r0 + 2, tr_ + 1):
        c = ws.cell(row=x, column=9, value=f'=IF($H${tr_}=0,0,H{x}/$H${tr_})'); c.number_format = PCT; c.font = font(bold=(x == tr_))
        c = ws.cell(row=x, column=10, value=f'=IF(B{x}=0,0,H{x}/B{x})'); c.number_format = USD; c.font = font(bold=(x == tr_))
        if x == tr_: ws.cell(row=x, column=9).fill = TOT; ws.cell(row=x, column=10).fill = TOT
    return tr_ + 2

r = 4
r = block(r, 'A. By category of student', CATNAMES, A_)
r = block(r, 'B. By student geography (cost-of-attendance rate applied)', GEOR, D_)
r = block(r, 'C. By enrollment geography', GEOS, C_)
r = block(r, 'D. By level', LEVELS, B_)
ws.cell(row=r, column=1, value='E. Spending-category totals (I-O model rows)').font = font(True, '003057', 11)
header(ws, r + 1, ['Spending category', 'Total ($)', 'Share', 'Suggested I-O treatment'])
IO = {'Books & supplies': 'Retail — book stores / educational supplies (retail margin only)',
      'Housing': 'Real estate — rental housing (on-campus portion = GT Housing revenue; see Assumptions)',
      'Food': 'Food service & drinking places / food & beverage stores (meal-plan portion is GT Dining revenue)',
      'Personal / misc.': 'Household consumption — distribute by a student/consumer spending pattern',
      'Transportation': 'Transit, gasoline (retail margin), air transportation for out-of-state travel'}
rr = r + 2
for k, it in enumerate(SPEND):
    ws.cell(row=rr, column=1, value=it).font = font()
    c = ws.cell(row=rr, column=2, value=f'=Allocation!{L(7+k)}{AL_TOT}'); c.number_format = USD; c.font = font(color=GREEN)
    c = ws.cell(row=rr, column=3, value=f'=IF($B${rr - k + len(SPEND)}=0,0,B{rr}/$B${rr - k + len(SPEND)})'); c.number_format = PCT; c.font = font()
    ws.cell(row=rr, column=4, value=IO[it]).font = font(size=9, color='555555')
    rr += 1
ws.cell(row=rr, column=1, value='Total').font = font(True)
c = ws.cell(row=rr, column=2, value=f'=SUM(B{r+2}:B{rr-1})'); c.number_format = USD; c.font = font(True); c.fill = TOT
c = ws.cell(row=rr, column=3, value=f'=SUM(C{r+2}:C{rr-1})'); c.number_format = PCT; c.font = font(True); c.fill = TOT
ws.column_dimensions['A'].width = 52; ws.column_dimensions['B'].width = 15
for cidx in range(3, 9): ws.column_dimensions[L(cidx)].width = 15
ws.column_dimensions['I'].width = 12; ws.column_dimensions['J'].width = 12; ws.column_dimensions['D'].width = 15
ws.freeze_panes = 'A4'
wb.save(OUT); print('saved', OUT)

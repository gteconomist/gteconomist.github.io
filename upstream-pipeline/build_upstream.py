"""Rebuild the upstream (traditional) IMPACT constant in index.html from CEDR's
FY impact workbook ("FY xx Economic Impact Emp Salary and ER Expense By Ga Region.xlsx").

Usage (from repo root):  python3 upstream-pipeline/build_upstream.py "upstream-pipeline/data/<workbook>.xlsx"

Reads the Data tab: regional rows 3-14, Georgia totals row 15, Final Table rows 26-32
(direct row 26 has Value Added and Output hard-coded from IMPLAN), state appropriation L43/M43.
Institute spending = Institute Spend + GTAA Spend (col U); indirect = Primary + Secondary (cols AH-AK).
Student spending impact block AS2:BD15 (rows 3-14 = regions 1-12, row 15 = Georgia): direct/indirect/induced
employment, labor income, value added and output. Direct student-spending jobs are part of the Final Table's
direct employment (I26 = GT employees + student-spending direct jobs).
"""
import openpyxl, json, re, sys
from pathlib import Path
XL=sys.argv[1]
SITE=Path(__file__).resolve().parent.parent/'index.html'
wb=openpyxl.load_workbook(XL,data_only=True); ws=wb['Data']
def c(r,col): 
    v=ws[f'{col}{r}'].value; return round(float(v),2) if isinstance(v,(int,float)) else v
emp_regions={}; inst_regions={}
for r in range(3,15):
    n=int(ws[f'A{r}'].value); name=ws[f'B{r}'].value
    emp_regions[str(n)]=dict(region=n,name=name,wages=c(r,'C'),employer_expense=c(r,'D'),employees=int(ws[f'E{r}'].value),
        direct_income=c(r,'H'),induced_income=c(r,'I'),total_income=c(r,'J'),induced_employment=c(r,'K'),value_added=c(r,'R'),output=c(r,'S'))
    inst_regions[str(n)]=dict(region=n,name=name,inst_spend=c(r,'U'),institute_spend=c(r,'F'),gtaa_spend=c(r,'G'),
        indirect_output=c(r,'AH'),indirect_emp=c(r,'AI'),indirect_income=c(r,'AJ'),
        induced_emp=c(r,'AD'),induced_output=c(r,'AC'),induced_income=c(r,'AE'))
r=15
emp_state=dict(wages=c(r,'C'),employer_expense=c(r,'D'),employees=int(ws['E15'].value),direct_income=c(r,'H'),induced_income=c(r,'I'),
    total_income=c(r,'J'),induced_employment=c(r,'K'),value_added=c(r,'R'),output=c(r,'S'))
inst_state=dict(inst_spend=c(r,'U'),institute_spend=c(r,'F'),gtaa_spend=c(r,'G'),indirect_output=c(r,'AH'),indirect_emp=c(r,'AI'),indirect_income=c(r,'AJ'),
    induced_emp=c(r,'AD'),induced_output=c(r,'AC'),induced_income=c(r,'AE'))
# student spending block AS:BD
SCOLS=dict(direct_emp='AS',indirect_emp='AT',induced_emp='AU',direct_income='AV',indirect_income='AW',induced_income='AX',
    direct_va='AY',indirect_va='AZ',induced_va='BA',direct_output='BB',indirect_output='BC',induced_output='BD')
def srow(r):
    d={k:c(r,col) for k,col in SCOLS.items()}
    for m in ['emp','income','va','output']:
        d['total_'+m]=round(d['direct_'+m]+d['indirect_'+m]+d['induced_'+m],2)
    return d
stu_regions={}
for r in range(3,15):
    n=int(ws[f'A{r}'].value); stu_regions[str(n)]=dict(region=n,name=ws[f'B{r}'].value,**srow(r))
stu_state=srow(15)
for k in ['total_emp','total_output']:
    assert abs(sum(v[k] for v in stu_regions.values())-stu_state[k])<1,k
def row(rr): return dict(employment=c(rr,'I'),labor_income=c(rr,'J'),value_added=c(rr,'K'),output=c(rr,'L'))
combined=dict(direct=row(26),indirect=row(28),induced=row(30),total=row(32))
# State appropriation share = appropriations / Georgia Tech revenues (direct output EXCLUDING student spending).
# The net adjustment applies to employee and Institute impact only, never to student spending.
approp=float(ws['L43'].value); gt_revenue=combined['direct']['output']-stu_state['direct_output']
share=round(approp/gt_revenue,8)
print(f"state approp share: {share:.4%} (workbook M43 = {float(ws['M43'].value):.4%}; GT revenue {gt_revenue:,.0f})")
IMPACT=dict(fiscal_year='FY25',state_approp_share=share,state_approp_amount=approp,
    combined=combined,employees=dict(statewide=emp_state,regions=emp_regions),institute=dict(statewide=inst_state,regions=inst_regions),
    students=dict(statewide=stu_state,regions=stu_regions,included_in_headline=True))
# sanity
for k in ['employment','labor_income','value_added','output']:
    s=sum(combined[t][k] for t in ['direct','indirect','induced']); assert abs(s-combined['total'][k])<1,(k,s)
assert abs(sum(v['indirect_output'] for v in inst_regions.values())-inst_state['indirect_output'])<1
print(json.dumps(combined,indent=1)); print(inst_state); print('students',stu_state)
assert abs(emp_state['employees']+stu_state['direct_emp']-combined['direct']['employment'])<1, 'direct employment != employees + student direct jobs'
# Final Table direct row should equal employee direct + student direct for every column (student spending is in the headline)
assert abs(emp_state['direct_income']+stu_state['direct_income']-combined['direct']['labor_income'])<1, \
    f"direct labor income {combined['direct']['labor_income']:,.0f} != employees {emp_state['direct_income']:,.0f} + students {stu_state['direct_income']:,.0f}"
print(f"direct value added {combined['direct']['value_added']:,.0f} (should include student direct VA {stu_state['direct_va']:,.0f})")
html=SITE.read_text()
new='const IMPACT = '+json.dumps(IMPACT)+';'
html2,n=re.subn(r'const IMPACT = \{.*?\};\n',lambda m:new+'\n',html,count=1,flags=re.S)
assert n==1
SITE.write_text(html2); print('updated',SITE)
